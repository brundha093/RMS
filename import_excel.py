from __future__ import annotations

import json
import sys
import urllib.parse
from datetime import datetime
from pathlib import Path

from database import RentalDatabase
from online_database import OnlineRentalDatabase


def normalize(value: object) -> str:
    return "_".join(str(value or "").strip().lower().replace("\n", " ").split())


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def date_text(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return text(value)


def import_workbook(workbook_path: str, database_path: str = "rental_manager.db") -> int:
    """Import flat rows from an .xlsx file using common Telugu/English headers."""
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("Install openpyxl before importing Excel data") from error

    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = list(sheet.values)
    if not rows:
        return 0
    headers = [normalize(value) for value in rows[0]]
    aliases = {
        "flat_number": ("flat", "flat_no", "flat_number", "ఫ్లాట్"),
        "flat_type": ("type", "flat_type", "రకం"),
        "tenant_name": ("tenant", "tenant_name", "name", "name_of_the_tenant", "పేరు"),
        "phone": ("phone", "mobile", "mobile_number", "phone_number", "ఫోన్"),
        "electricity_number": ("electricity_number", "electricity_no", "service_number", "electricity_bill_number", "కరెంట్_నంబర్"),
        "monthly_rent": ("rent", "monthly_rent", "rent_amount", "అద్దె"),
        "maintenance": ("maintenance", "maintenance_charge", "maintanence_charge", "మెయింటెనెన్స్"),
        "parking": ("parking", "car_parking", "car_parking_charge", "పార్కింగ్"),
        "bike_charge": ("electric_bike", "electric_bike_charge", "bike_charge"),
        "repair": ("repair", "repairs", "repair_charge"),
        "joining_date": ("date_of_joining", "date_of_jioning", "joining_date"),
        "vacating_date": ("date_of_exit", "exit_date", "vacating_date"),
        "advance_amount": ("advance", "advance_amount", "advance_amount_paid", "deposit", "security_deposit"),
    }
    indexes = {}
    for target, names in aliases.items():
        indexes[target] = next((headers.index(name) for name in names if name in headers), None)

    database = RentalDatabase(database_path)
    imported = 0
    imported_flat_numbers: set[str] = set()
    for row in rows[1:]:
        flat_number = row[indexes["flat_number"]] if indexes["flat_number"] is not None else None
        if not flat_number:
            continue
        values = {key: row[index] if index is not None and index < len(row) else "" for key, index in indexes.items()}
        database.connection.execute(
            """INSERT INTO flats (flat_number, flat_type, tenant_name, phone, electricity_number, monthly_rent,
            maintenance, parking, repair, bike_charge, joining_date, vacating_date, advance_amount)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(flat_number) DO UPDATE SET tenant_name=excluded.tenant_name, phone=excluded.phone,
            electricity_number=excluded.electricity_number, monthly_rent=excluded.monthly_rent,
            flat_type=excluded.flat_type, maintenance=excluded.maintenance, parking=excluded.parking,
            repair=excluded.repair, bike_charge=excluded.bike_charge, joining_date=excluded.joining_date,
            vacating_date=excluded.vacating_date, advance_amount=excluded.advance_amount""",
            (text(values["flat_number"]), text(values["flat_type"]), text(values["tenant_name"]), text(values["phone"]),
             text(values["electricity_number"]), float(values["monthly_rent"] or 0), float(values["maintenance"] or 0),
             float(values["parking"] or 0), float(values["repair"] or 0), float(values["bike_charge"] or 0),
             date_text(values["joining_date"]), date_text(values["vacating_date"]), float(values["advance_amount"] or 0)),
        )
        imported_flat_numbers.add(text(values["flat_number"]))
        imported += 1
    if imported_flat_numbers:
        placeholders = ",".join("?" for _ in imported_flat_numbers)
        obsolete_ids = database.connection.execute(
            f"SELECT id FROM flats WHERE flat_number NOT IN ({placeholders})",
            tuple(imported_flat_numbers),
        ).fetchall()
        for (flat_id,) in obsolete_ids:
            database.connection.execute("DELETE FROM payments WHERE flat_id = ?", (flat_id,))
        database.connection.execute(
            f"DELETE FROM flats WHERE flat_number NOT IN ({placeholders})",
            tuple(imported_flat_numbers),
        )
    database.connection.commit()
    database.close()
    return imported


def normalize_electricity_number(value: object) -> str:
    text_value = str(value or "").strip()
    if text_value.endswith(".0"):
        text_value = text_value[:-2]
    return text_value


def load_electricity_bill_updates(json_path: str | Path = "electricity_bills.json") -> dict[str, dict[str, float]]:
    resolved = Path(json_path)
    if not resolved.is_absolute():
        resolved = Path(__file__).resolve().parent / resolved
    if not resolved.exists():
        return {}
    payload = json.loads(resolved.read_text(encoding="utf-8"))
    updates: dict[str, dict[str, float]] = {}
    for entry in payload:
        electricity_number = normalize_electricity_number(entry.get("electricity_number"))
        if not electricity_number:
            continue
        total_amount = float(str(entry.get("total_amount") or 0).replace(",", "").strip() or 0)
        arrears = float(str(entry.get("arrears") or 0).replace(",", "").strip() or 0)
        updates[electricity_number] = {"total_amount": total_amount, "arrears": arrears}
    return updates


def update_excel_electricity_totals(workbook_path: str, json_path: str = "electricity_bills.json") -> int:
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError("Install openpyxl before updating electricity totals") from error

    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=False, min_row=1))
    if not rows:
        workbook.close()
        return 0

    header_row = [normalize(value.value) if value.value is not None else "" for value in rows[0]]
    indexes = {
        "electricity_number": next((idx for idx, name in enumerate(header_row) if name in {"electricity_bill_number", "electricity_number", "service_number", "కరెంట్_నంబర్"}), None),
        "electricity_total": next((idx for idx, name in enumerate(header_row) if name in {"electricity_total", "electricity_total_amount", "total_amount"}), None),
        "electricity_arrears": next((idx for idx, name in enumerate(header_row) if name in {"electricity_arrears", "arrears", "electricityarrears"}), None),
    }
    if indexes["electricity_number"] is None:
        workbook.close()
        return 0

    updates = load_electricity_bill_updates(json_path)
    updated = 0
    for row in rows[1:]:
        electricity_number = row[indexes["electricity_number"]].value if indexes["electricity_number"] < len(row) else None
        lookup_key = normalize_electricity_number(electricity_number)
        if lookup_key not in updates:
            continue
        values = updates[lookup_key]
        if indexes["electricity_total"] is not None and indexes["electricity_total"] < len(row):
            row[indexes["electricity_total"]].value = values["total_amount"]
        if indexes["electricity_arrears"] is not None and indexes["electricity_arrears"] < len(row):
            row[indexes["electricity_arrears"]].value = values["arrears"]
        updated += 1
    workbook.save(workbook_path)
    workbook.close()
    return updated


def sync_electricity_from_json(workbook_path: str = "Excel.xlsx", json_path: str = "electricity_bills.json", database_path: str = "rental_manager.db") -> dict[str, int]:
    excel_updated = update_excel_electricity_totals(workbook_path, json_path)
    database = RentalDatabase(database_path)
    updates = load_electricity_bill_updates(json_path)
    synced = 0
    for electricity_number, values in updates.items():
        row = database.connection.execute(
            "SELECT id, electricity_due, electricity_paid, electricity_arrears FROM flats WHERE electricity_number = ?",
            (electricity_number,),
        ).fetchone()
        if row is None:
            continue
        due = float(values["total_amount"])
        paid = max(due - float(values["arrears"]), 0)
        arrears = float(values["arrears"])
        database.connection.execute(
            "UPDATE flats SET electricity_due = ?, electricity_paid = ?, electricity_arrears = ? WHERE id = ?",
            (due, paid, arrears, row["id"]),
        )
        synced += 1
    database.connection.commit()
    try:
        cloud = OnlineRentalDatabase()
        for electricity_number, values in updates.items():
            row = cloud._request("flats", query=f"?electricity_number=eq.{urllib.parse.quote_plus(electricity_number)}&select=id,electricity_number,electricity_due,electricity_paid,electricity_arrears")
            if not row:
                continue
            flat_id = row[0]["id"]
            due = float(values["total_amount"])
            paid = max(due - float(values["arrears"]), 0)
            cloud.update_electricity(flat_id, due, paid, float(values["arrears"]))
    except Exception:
        pass
    database.close()
    return {"excel_rows_updated": excel_updated, "sqlite_rows_updated": synced}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python import_excel.py Excel.xlsx")
    print(f"Imported {import_workbook(sys.argv[1])} flat records")
    sync_result = sync_electricity_from_json(sys.argv[1], "electricity_bills.json", "rental_manager.db")
    print(f"Updated electricity data in Excel and database: {sync_result}")
